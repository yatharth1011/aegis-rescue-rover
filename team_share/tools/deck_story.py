"""
Story slides (3 onwards) for the A.E.G.I.S. proposal deck. Called by build_submission.py.
Every number here comes from the project's own results files; nothing is invented.
"""
import json
import os

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

NAVY = RGBColor(0x0B, 0x12, 0x20)
NAVY2 = RGBColor(0x15, 0x1F, 0x33)
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x47, 0x55, 0x69)
SOFT = RGBColor(0xA9, 0xB4, 0xC6)
ACCENT = RGBColor(0xEA, 0x6C, 0x1C)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
PANEL = RGBColor(0xF1, 0xF5, 0xF9)
LINE = RGBColor(0xCB, 0xD5, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = Emu(12192000), Emu(6858000)
M = Emu(457200)                       # 0.5 in margin
IN = 914400


def E(inches):
    return Emu(int(inches * IN))


def bg(slide, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def box(slide, x, y, w, h, fill, line=None, radius=0.06):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def txt(slide, x, y, w, h, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.08):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if not isinstance(paras, list):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if i:
            p.space_before = Pt(5)
        runs = para if isinstance(para, list) else [(para, {})]
        for t, st in runs:
            r = p.add_run()
            r.text = t
            r.font.name = "Calibri"
            r.font.size = Pt(st.get("size", size))
            r.font.bold = st.get("bold", bold)
            r.font.color.rgb = st.get("color", color)
    return tb


def pic_fit(slide, path, x, y, w, h, align="center"):
    from PIL import Image
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = int(iw * s), int(ih * s)
    px = x + (w - pw) // 2 if align == "center" else x
    return slide.shapes.add_picture(path, px, y + (h - ph) // 2, pw, ph)


def title(slide, kicker, head, dark=True):
    txt(slide, M, E(0.38), W - 2 * M, E(0.3), kicker.upper(), size=12, bold=True, color=ACCENT)
    txt(slide, M, E(0.66), W - 2 * M, E(0.7), head, size=30, bold=True, color=WHITE if dark else INK)


def footer(slide, n, dark=True):
    txt(slide, M, H - E(0.42), W - 2 * M, E(0.25),
        [[("A.E.G.I.S.  |  TL101 Tinkering Lab, IIT Ropar  |  ", {}), ("%d" % n, {"bold": True})]],
        size=10, color=SOFT if dark else MUTED)


def stat(slide, x, y, w, big, small, dark=True, big_color=None):
    txt(slide, x, y, w, E(0.6), big, size=34, bold=True, color=big_color or (WHITE if dark else INK))
    txt(slide, x, y + E(0.62), w, E(0.6), small, size=12, color=SOFT if dark else MUTED)


def add_story(prs, paths, sim):
    blank = prs.slide_layouts[6]
    n = 3

    # ---------- 3. the mission ---------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    title(s, "The idea", "Hear.  Search.  Recognise.  Rescue.  Return.")
    steps = [("01", "Hear", "Two microphones find the direction of a scream (GCC-PHAT, 48 kHz)."),
             ("02", "Search", "Laser sensors map the table; the rover inspects every unseen spot."),
             ("03", "Recognise", "A neural network (YOLO) spots the victim in the camera feed."),
             ("04", "Rescue", "The rover drives the victim into its bay; a door closes behind it."),
             ("05", "Return", "Shortest path home over ground it has already proven safe.")]
    cw, gap = E(2.24), E(0.15)
    for i, (num, head, body) in enumerate(steps):
        x = M + i * (cw + gap)
        box(s, x, E(1.75), cw, E(2.75), NAVY2)
        txt(s, x + E(0.2), E(1.95), cw - E(0.4), E(0.5), num, size=26, bold=True, color=ACCENT)
        txt(s, x + E(0.2), E(2.55), cw - E(0.4), E(0.45), head, size=20, bold=True, color=WHITE)
        txt(s, x + E(0.2), E(3.05), cw - E(0.4), E(1.4), body, size=13, color=SOFT)
    box(s, M, E(4.8), W - 2 * M, E(1.25), NAVY2)
    txt(s, M + E(0.3), E(4.98), W - 2 * M - E(0.6), E(1.0),
        [[("Two live demos on a lab table.  ", {"bold": True, "color": WHITE}),
          ("Run 1: the victim is hidden behind books and there is no sound, so the rover must systematically inspect the "
           "whole table.  Run 2: a speaker at the victim plays a scream and the rover homes in on the sound.", {})],
         [("A low-cost prototype of the search-and-rescue loop used for collapsed buildings and wildfire zones.", {"color": SOFT})]],
        size=15, color=WHITE)
    footer(s, n)
    n += 1

    # ---------- 4. meet Mk2 ------------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    title(s, "Designed in Autodesk Fusion", "Meet A.E.G.I.S. Mk2")
    pic_fit(s, paths["render_open"], M, E(1.45), E(7.4), E(4.75))
    specs = [("41", "components, every purchased part modelled; 0 clashes"),
             ("64 x 46 mm", "intake bay with a servo-driven door"),
             ("4WD", "N20 150:1 gearmotors with encoders"),
             ("2 brains", "ESP32-S3 on board, laptop over Wi-Fi"),
             ("13 sensors", "3 laser ToF, 2 mics, IMU, camera, 6 IR"),
             ("263 x 140 x 148 mm", "about 1 kg, 3S Li-ion 10C pack")]
    y = E(1.55)
    for big, small in specs:
        txt(s, E(8.1), y, E(4.6), E(0.4), big, size=20, bold=True, color=ACCENT)
        txt(s, E(8.1), y + E(0.36), E(4.6), E(0.4), small, size=13, color=SOFT)
        y += E(0.78)
    footer(s, n)
    n += 1

    # ---------- 5. every part, every wire ----------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Build-ready", "Every part placed, every wire routed", dark=False)
    pic_fit(s, paths["render_wiring"], M, E(1.45), E(5.0), E(3.6))
    pic_fit(s, paths["wiring_guide"], E(5.65), E(1.45), E(3.8), E(3.6))
    pic_fit(s, paths["render_closed"], E(9.55), E(1.45), E(3.3), E(3.6))
    caps = [(M, E(5.0), "Wiring harness, colour-coded"), (E(5.65), E(3.8), "Pin-exact wiring guide (PDF)"),
            (E(9.55), E(3.3), "Door closed: victim inside")]
    for x, w, c in caps:
        txt(s, x, E(5.1), w, E(0.3), c, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    box(s, M, E(5.5), W - 2 * M, E(0.75), PANEL, LINE)
    txt(s, M + E(0.25), E(5.62), W - 2 * M - E(0.5), E(0.55),
        [[("Ready to fabricate: ", {"bold": True}),
          ("laser-cut and engrave DXFs (placement marks on both decks), STLs for 4 printed parts, STEP model, "
           "a 35-line BOM with live Robu links (wiring included), and firmware that builds for the ESP32-S3.", {})]], size=14)
    footer(s, n, dark=False)
    n += 1

    # ---------- 6. architecture ------------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    title(s, "Split brain, all our own code", "Reflexes on the robot. Strategy on the laptop.")
    cols = [("ON THE ROBOT", "ESP32-S3 + ESP32-CAM",
             ["Wheel speed control (PI) and odometry", "Speed limit near obstacles (control barrier function)",
              "Table-edge reflex, stall escape", "Stops if Wi-Fi goes quiet for 0.3 s", "Scream direction (GCC-PHAT, 48 kHz)",
              "Camera stream, telemetry"]),
            ("OVER WI-FI", "UDP, about 2 Mbit/s",
             ["Camera frames (JPEG, 8 fps)", "Sensor telemetry (20 Hz)", "Audio for scream detection",
              "Motion commands (10 Hz)", "Tested with 20 % packet loss and 3 s drop-outs"]),
            ("ON THE LAPTOP", "The edge brain (Mac M5)",
             ["YOLO victim detector", "Scream detector", "Live map of the table", "Coverage search: inspect everything",
              "A* shortest path home", "Trains the AI models (GPU)"])]
    cw = E(3.85)
    for i, (k, sub, items) in enumerate(cols):
        x = M + i * (cw + E(0.2))
        box(s, x, E(1.65), cw, E(4.45), NAVY2 if i != 1 else RGBColor(0x1F, 0x2A, 0x40))
        txt(s, x + E(0.25), E(1.85), cw - E(0.5), E(0.3), k, size=12, bold=True, color=ACCENT)
        txt(s, x + E(0.25), E(2.15), cw - E(0.5), E(0.45), sub, size=18, bold=True, color=WHITE)
        txt(s, x + E(0.25), E(2.75), cw - E(0.5), E(3.2), ["•  " + t for t in items], size=13, color=SOFT)
    footer(s, n)
    n += 1

    # ---------- 7-8. the GIFs ----------------------------------------------------------------------------
    for gif, kick, head, facts in (
            (paths["gif1"], "3D simulation on a Mac (physics, camera, sound, Wi-Fi)  |  demo 1", "Hidden victim: inspect everything, rescue, return",
             [("%.0f s" % paths["ex1"]["first_seen"], "victim spotted"), ("%.0f s" % paths["ex1"]["captured"], "victim inside the bay"),
              ("%.0f s" % paths["ex1"]["t_end"], "back home, door shut")]),
            (paths["gif2"], "3D simulation on a Mac (physics, camera, sound, Wi-Fi)  |  demo 2", "Scream homing: hear it, find it, bring it back",
             [("%.0f s" % paths["ex2"]["first_seen"], "victim spotted"), ("%.0f s" % paths["ex2"]["captured"], "victim inside the bay"),
              ("%.0f s" % paths["ex2"]["t_end"], "back home, door shut")])):
        s = prs.slides.add_slide(blank)
        bg(s, NAVY)
        title(s, kick, head)
        pic_fit(s, gif, M, E(1.45), W - 2 * M, E(4.35))
        for i, (big, small) in enumerate(facts):
            x = M + E(0.6) + i * E(4.0)
            txt(s, x, E(5.85), E(1.3), E(0.5), big, size=24, bold=True, color=ACCENT)
            txt(s, x + E(1.15), E(5.95), E(1.8), E(0.4), small, size=13, color=SOFT)
        footer(s, n)
        n += 1

    # ---------- 9. the AI we trained -------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Trained on our own GPU", "What the robot's AI sees", dark=False)
    pic_fit(s, paths["ai_sees"], M, E(1.45), E(7.6), E(4.2))
    txt(s, M, E(5.7), E(7.6), E(0.3), "Examples from held-out simulated ESP32-CAM frames; look-alike decoys are present in the scenes.",
        size=11, color=MUTED)
    y = E(1.5)
    for big, small in (("98.9 %", "of victims found (recall) on 800 frames it never trained on; 98.7 % precision"),
                       ("4,800", "rendered camera frames, labelled automatically; trained on the Mac GPU in 37 min"),
                       ("80 / 80", "screams detected, 0 false alarms on 160 talking / fan clips (synthetic sounds)"),
                       ("8°", "median scream-direction error, after resolving front/back by turning")):
        txt(s, E(8.4), y, E(4.4), E(0.5), big, size=28, bold=True, color=ACCENT)
        txt(s, E(8.4), y + E(0.5), E(4.4), E(0.6), small, size=12, color=MUTED)
        y += E(1.15)
    footer(s, n, dark=False)
    n += 1

    # ---------- 10. results ----------------------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "120 simulated missions", "It works, and we measured how well", dark=False)
    sets = sim.get("sets", {})
    nums = [("%d %%" % round(sets["final_demo1"]["success_pct"]), "hidden-victim missions\ncompleted end to end"),
            ("%d %%" % round(sets["final_demo2"]["success_pct"]), "scream-homing missions\ncompleted end to end"),
            ("%d %%" % round((sets["final_demo1"]["found_pct"] + sets["final_demo2"]["found_pct"]) / 2), "of missions where the victim\nwas found (both demos)"),
            ("%d-%d %%" % tuple(sorted((round(sets["final_demo1_badlink"]["success_pct"]), round(sets["final_demo2_badlink"]["success_pct"])))),
             "still completed on a bad\nWi-Fi link (20 % loss)")]
    for i, (big, small) in enumerate(nums):
        x = M + i * E(3.1)
        box(s, x, E(1.5), E(2.95), E(1.45), PANEL, LINE)
        txt(s, x + E(0.2), E(1.6), E(2.6), E(0.6), big, size=34, bold=True, color=ACCENT)
        txt(s, x + E(0.2), E(2.25), E(2.6), E(0.65), small, size=12, color=MUTED)
    pic_fit(s, paths["summary"], M, E(3.1), W - 2 * M, E(3.2))
    footer(s, n, dark=False)
    n += 1

    # ---------- 11. engineering rigour ---------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    title(s, "Engineering, not guesswork", "Checked before a single part is bought")
    tiles = [("Firmware", "Builds for the ESP32-S3 (ESP-IDF 5.3, FreeRTOS on both cores). 32/32 unit tests pass."),
             ("Power", "Worst case 5.8 A = 2.7 C on a 10 C pack. Fused and switched; separate 5 V rails for logic and servo."),
             ("Drive", "Hardest manoeuvre needs 27 % of motor torque. 0.37 m/s top speed."),
             ("Terrain (3D physics)", "Climbs 25° slopes and 1 cm bumps; a 6 mm kerb stops it. Built for tables and floors."),
             ("Simulation found bugs", "Bay trigger, door clearance, front/back sound confusion, zig-zag search, pinned recovery, rear-corner falls: all fixed."),
             ("Budget", "Rs 11,971 of Rs 15,000, 35 items including all wiring, 34 with live Robu.in links checked.")]
    tw, th = E(3.85), E(2.1)
    for i, (head, body) in enumerate(tiles):
        x = M + (i % 3) * (tw + E(0.2))
        y = E(1.6) + (i // 3) * (th + E(0.2))
        box(s, x, y, tw, th, NAVY2)
        txt(s, x + E(0.25), y + E(0.2), tw - E(0.5), E(0.4), head, size=17, bold=True, color=ACCENT)
        txt(s, x + E(0.25), y + E(0.68), tw - E(0.5), th - E(0.8), body, size=13, color=SOFT)
    footer(s, n)
    n += 1

    # ---------- 12. budget ---------------------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Budget", "Rs 11,971 of Rs 15,000, with Rs 3,029 in reserve", dark=False)
    groups = [("Motion: 4 gearmotors, 2 drivers, wheels", 2690), ("Power: 3S pack, charger, 2 buck converters", 2576),
              ("Brains and vision: ESP32-S3, ESP32-CAM, programmer", 1997), ("Sensing: ToF x3, IMU, mics, IR x6, I/O expander", 1782),
              ("Structure: PLA+, acrylic, M3 hardware", 1514), ("Rescue mechanism: MG946R servo", 459),
              ("Wiring: wire, connectors, switch, fuse, carrier board", 953)]
    total_w = W - 2 * M
    x = M
    cols_ = [RGBColor(0xEA, 0x6C, 0x1C), RGBColor(0x1F, 0x4E, 0x99), RGBColor(0x0F, 0x76, 0x6E), RGBColor(0x7C, 0x3A, 0xED),
             RGBColor(0x64, 0x74, 0x8B), RGBColor(0xDC, 0x26, 0x26), RGBColor(0xCA, 0x8A, 0x04)]
    for (name, v), c in zip(groups, cols_):
        w = int(total_w * v / 15000)
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, E(1.7), w, E(0.7))
        r.fill.solid()
        r.fill.fore_color.rgb = c
        r.line.fill.background()
        x += w
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, E(1.7), M + total_w - x, E(0.7))
    r.fill.solid()
    r.fill.fore_color.rgb = PANEL
    r.line.color.rgb = LINE
    txt(s, x, E(1.85), M + total_w - x, E(0.4), "reserve", size=13, color=MUTED, align=PP_ALIGN.CENTER)
    for i, ((name, v), c) in enumerate(zip(groups, cols_)):
        y = E(2.6) + i * E(0.47)
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, M, y + E(0.07), E(0.22), E(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = c
        dot.line.fill.background()
        txt(s, M + E(0.4), y, E(8.5), E(0.4), name, size=15)
        txt(s, E(9.5), y, E(1.5), E(0.4), "Rs %s" % format(v, ","), size=15, bold=True, align=PP_ALIGN.RIGHT)
    txt(s, M, E(5.9), W - 2 * M, E(0.4), "Prices from Robu.in (acrylic sheet from Amazon.in, not sold on Robu), 23-25 Oct 2026. The laptop brain is the team's existing Mac: no Jetson needed.",
        size=12, color=MUTED)
    footer(s, n, dark=False)
    n += 1

    # ---------- 13. honest limits and plan -----------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "What we know is not solved yet", "Honest limits, and the plan to close them", dark=False)
    limits = ["Edge safety costs search time: rear-corner edge sensors cut falls to 1 in 120 missions (on bad Wi-Fi), but the wider edge margin lowered hidden-victim success from 80 % to 70 %; tune it",
              "Detector trained on renders: retrain on real photos of our victim token",
              "Synthetic screams: record real ones for the scream detector",
              "Flat surfaces only: rugged terrain needs bigger wheels and suspension (future Mk3)"]
    txt(s, M, E(1.55), E(6.0), E(4.4), [[("•  ", {"color": ACCENT, "bold": True}), (t, {})] for t in limits], size=15)
    plan = [("Wk 1", "Order parts, laser-cut and print"), ("Wk 2", "Wire, flash firmware, bring-up"),
            ("Wk 3", "Real camera + scream data, retrain"), ("Wk 4", "Laptop brain on the real robot"),
            ("Wk 5", "Tune, 10+ trial runs per demo"), ("Wk 6", "Demo and report")]
    for i, (wk, what) in enumerate(plan):
        y = E(1.55) + i * E(0.72)
        box(s, E(6.9), y, E(1.0), E(0.55), ACCENT)
        txt(s, E(6.9), y + E(0.12), E(1.0), E(0.35), wk, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        txt(s, E(8.05), y + E(0.12), E(4.7), E(0.4), what, size=15)
    footer(s, n, dark=False)
    n += 1

    # ---------- 14. team ---------------------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    title(s, "The team", "Six AI & Data Engineering students, one rover")
    team = [("Yatharth Shrivastava", "", ""), ("Prabhnoor Singh", "", ""),
            ("Shreyansh", "", ""), ("Souvik Das", "", ""),
            ("Utkarsh Kumar Gautam", "", ""), ("Shubham Maurya", "", "")]
    cw, ch = E(3.85), E(1.15)
    for i, (name, roll, role) in enumerate(team):
        x = M + (i % 3) * (cw + E(0.2))
        y = E(1.6) + (i // 3) * (ch + E(0.2))
        box(s, x, y, cw, ch, NAVY2)
        txt(s, x + E(0.25), y + E(0.2), cw - E(0.5), E(0.4), name, size=18, bold=True, color=WHITE)
        txt(s, x + E(0.25), y + E(0.62), cw - E(0.5), E(0.4), roll + ("   |   " + role if role else ""), size=13, color=SOFT)
    txt(s, M, E(4.45), W - 2 * M, E(0.35), "WORK PACKAGES  (one owner each, assigned at kick-off)", size=12, bold=True, color=ACCENT)
    wps = ["Vision and YOLO", "Audio and scream direction", "Firmware and control", "CAD and fabrication",
           "Electronics and power", "Simulation and RL"]
    for i, w_ in enumerate(wps):
        x = M + (i % 3) * (cw + E(0.2))
        y = E(4.85) + (i // 3) * E(0.55)
        txt(s, x, y, cw, E(0.4), [[("■  ", {"color": ACCENT}), (w_, {})]], size=15, color=WHITE)
    footer(s, n)
    n += 1

    # ---------- 15. the ask -------------------------------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    txt(s, M, E(1.6), W - 2 * M, E(0.4), "OUR REQUEST", size=14, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    txt(s, M, E(2.1), W - 2 * M, E(1.2), "Approve A.E.G.I.S. for the TL101 build", size=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s, E(1.5), E(3.5), W - E(3.0), E(1.6),
        [[("Designed in Fusion, simulated end to end, budgeted at Rs 11,971. ", {"bold": True, "color": WHITE}),
          ("On demo day the rover will find a hidden victim, home in on a scream, take the victim inside "
           "and bring it back, on a real table, in front of you.", {})]],
        size=18, color=SOFT, align=PP_ALIGN.CENTER)
    txt(s, M, E(5.6), W - 2 * M, E(0.4), "Hear.  Search.  Recognise.  Rescue.  Return.", size=16, bold=True, color=ACCENT,
        align=PP_ALIGN.CENTER)
    footer(s, n)
    return n
